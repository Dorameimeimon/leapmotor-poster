#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 agent-browser 拉取的分页 JSON 合并生成晨夕会审核列表 xlsx（与手工导出同构）。
输入: /tmp/leap_p{1..6}.raw  (agent-browser eval 输出, JSON编码字符串)
输出: workspace/晨夕会审核列表<date>_web.xlsx  (B=晨夕会类型, D=门店, G=会议日期)
"""
import json, datetime, sys, glob
from openpyxl import Workbook

WS = '/Users/mei/WorkBuddy/2026-06-23-10-23-34'
today = sys.argv[1] if len(sys.argv) > 1 else 'latest'

recs = {}
for f in sorted(glob.glob('/tmp/leap_p*.raw')):
    try:
        raw = open(f).read().strip()
        s = json.loads(raw)          # CLI 输出是 JSON 编码的字符串
        d = json.loads(s)            # 内层才是接口响应
    except Exception:
        continue
    if isinstance(d, dict) and d.get('code') == 200:
        for r in d['data']['records']:
            recs[r['checkId']] = r   # checkId 去重

if not recs:
    print('ERROR: no records fetched'); sys.exit(1)

times = sorted(r['uploadTime'] for r in recs.values())
wb = Workbook(); ws = wb.active
ws.append(['序号', '晨夕会类型', '备注', '门店', '大区', '区域', '会议日期'])
for idx, r in enumerate(sorted(recs.values(), key=lambda x: x['uploadTime']), start=2):
    dt = datetime.datetime.strptime(r['uploadTime'], '%Y-%m-%d %H:%M:%S')
    ws.cell(row=idx, column=2, value=r['meetingTypeLabel'])
    ws.cell(row=idx, column=4, value=r['storeName'])
    ws.cell(row=idx, column=7, value=dt)
out = f'{WS}/晨夕会审核列表{today}_web.xlsx'
wb.save(out)
print(f'xlsx saved: {out}')
print(f'records={len(recs)} span {times[0]} ~ {times[-1]}')
